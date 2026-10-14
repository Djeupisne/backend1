# ═══════════════════════════════════════════════════════
# Generation des rapports Excel / PDF / CSV / Word
# Contenu repris verbatim de l ancien server.py monolithique
# (refactor purement structurel — aucune logique modifiee)
# ═══════════════════════════════════════════════════════
import io, csv, json, datetime, re
from ..config import (logger,
                     REPORTLAB_AVAILABLE, OPENPYXL_AVAILABLE,
                     SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak, KeepTogether,
                     getSampleStyleSheet, ParagraphStyle, landscape, A4, colors,
                     cm, inch, TA_CENTER, TA_LEFT, TA_RIGHT, canvas,
                     Workbook, Font, PatternFill, Alignment, Border, Side, get_column_letter,
                     BarChart, Reference, PieChart, DataLabelList)

def generate_excel_report_enhanced(candidats, poste_filter=""):
    if not OPENPYXL_AVAILABLE:
        return None
    try:
        sorted_candidats = sort_candidats(candidats)
        wb = Workbook()
        
        # ===================== FEUILLE 1 : LISTE DES CANDIDATS =====================
        ws = wb.active
        ws.title = "Candidats"
        
        # Styles modernes
        header_font = Font(bold=True, size=11, color="FFFFFF", name='Segoe UI')
        header_fill = PatternFill(start_color="0d47a1", end_color="0d47a1", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        number_alignment = Alignment(horizontal="center", vertical="center")
        
        # Bordures modernes
        thin_border = Border(
            left=Side(style='thin', color='1565c0'),
            right=Side(style='thin', color='1565c0'),
            top=Side(style='thin', color='1565c0'),
            bottom=Side(style='thin', color='1565c0')
        )
        
        # En-têtes
        headers = [
            "Rang", "N° Dossier", "Nom", "Prénom", "Email", "Téléphone",
            "Poste", "Statut", "Score", "Décision",
            "Points Forts", "Points de Vigilance", "Synthèse", "Motifs"
        ]
        col_widths = [5, 12, 18, 18, 25, 15, 30, 14, 8, 22, 35, 35, 40, 35]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
            ws.column_dimensions[get_column_letter(col)].width = col_widths[col-1]
        
        # Lignes de données
        row_idx = 2
        rank = 1
        for c in sorted_candidats:
            score_breakdown = c.get('score_breakdown_parsed', {})
            analyse_details = c.get('analyse_details_parsed', {})
            points_forts = analyse_details.get('points_forts', []) or score_breakdown.get('points_forts', [])
            points_vigilance = analyse_details.get('points_vigilance', []) or score_breakdown.get('points_vigilance', [])
            synthese = analyse_details.get('synthese_recruteur', '') or c.get('synthese', '')
            statut = c.get('statut', 'en_attente')
            decision = c.get('decision', '')
            flags_elim = analyse_details.get('flags_eliminatoires', []) or score_breakdown.get('flags_eliminatoires', [])
            
            # Construction du motif
            motif_parts = []
            if flags_elim:
                motif_parts.append(f"CRITÈRES ÉLIMINATOIRES: {', '.join(flags_elim[:3])}")
            if points_vigilance:
                motif_parts.append(f"POINTS DE VIGILANCE: {', '.join(points_vigilance[:2])}")
            if points_forts:
                motif_parts.append(f"POINTS FORTS: {', '.join(points_forts[:2])}")
            if synthese:
                motif_parts.append(f"SYNTHÈSE: {synthese[:150]}")
            
            motif = " | ".join(motif_parts) if motif_parts else "Analyse en cours..."
            
            try:
                score = float(c.get('score', 0))
            except (ValueError, TypeError):
                score = 0.0
            
            # Écriture des données
            ws.cell(row=row_idx, column=1, value=rank).alignment = number_alignment
            ws.cell(row=row_idx, column=2, value=c.get('numero_dossier', '')).alignment = center_alignment
            ws.cell(row=row_idx, column=3, value=c.get('nom', '')).alignment = cell_alignment
            ws.cell(row=row_idx, column=4, value=c.get('prenom', '')).alignment = cell_alignment
            ws.cell(row=row_idx, column=5, value=c.get('email', '')).alignment = cell_alignment
            ws.cell(row=row_idx, column=6, value=c.get('telephone', '')).alignment = cell_alignment
            ws.cell(row=row_idx, column=7, value=c.get('poste', '')).alignment = cell_alignment
            ws.cell(row=row_idx, column=8, value=statut.upper()).alignment = center_alignment
            ws.cell(row=row_idx, column=9, value=score).alignment = center_alignment
            ws.cell(row=row_idx, column=10, value=decision).alignment = cell_alignment
            ws.cell(row=row_idx, column=11, value=", ".join(points_forts[:3]) if points_forts else '').alignment = cell_alignment
            ws.cell(row=row_idx, column=12, value=", ".join(points_vigilance[:3]) if points_vigilance else '').alignment = cell_alignment
            ws.cell(row=row_idx, column=13, value=synthese[:250] if synthese else '').alignment = cell_alignment
            ws.cell(row=row_idx, column=14, value=motif).alignment = cell_alignment
            
            # Couleur de fond selon statut
            if statut == 'retenu':
                fill_color = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")
            elif statut == 'entretien':
                fill_color = PatternFill(start_color="fff3e0", end_color="fff3e0", fill_type="solid")
            elif statut == 'rejete':
                fill_color = PatternFill(start_color="ffebee", end_color="ffebee", fill_type="solid")
            else:
                fill_color = PatternFill(start_color="f5f5f5", end_color="f5f5f5", fill_type="solid")
            
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = thin_border
                cell.fill = fill_color
            
            row_idx += 1
            rank += 1
        
        # Total
        total_row = row_idx
        if len(sorted_candidats) > 0:
            total_cell = ws.cell(row=total_row, column=1, value=f"Total: {len(sorted_candidats)} candidats")
            total_cell.font = Font(bold=True, size=12, color="0d47a1", name='Segoe UI')
            ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
            total_cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # ===================== FEUILLE 2 : STATISTIQUES =====================
        ws_stats = wb.create_sheet("Statistiques")
        
        # Titre
        title_cell = ws_stats.cell(row=1, column=1, value="📊 STATISTIQUES DES CANDIDATURES")
        title_cell.font = Font(bold=True, size=16, color="0d47a1", name='Segoe UI')
        ws_stats.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Statistiques générales
        stats_data = [
            ["Indicateur", "Valeur"],
            ["Total candidats", len(sorted_candidats)],
            ["Retenus", sum(1 for c in sorted_candidats if c.get('statut') == 'retenu')],
            ["Entretien", sum(1 for c in sorted_candidats if c.get('statut') == 'entretien')],
            ["En attente", sum(1 for c in sorted_candidats if c.get('statut') == 'en_attente')],
            ["Rejetés", sum(1 for c in sorted_candidats if c.get('statut') == 'rejete')],
            ["Score moyen", f"{sum(float(c.get('score', 0)) for c in sorted_candidats) / max(1, len(sorted_candidats)):.2f}"],
            ["Score max", max([float(c.get('score', 0)) for c in sorted_candidats]) if sorted_candidats else 0],
            ["Score min", min([float(c.get('score', 0)) for c in sorted_candidats]) if sorted_candidats else 0],
        ]
        
        for i, row_data in enumerate(stats_data, 3):
            for j, val in enumerate(row_data, 1):
                cell = ws_stats.cell(row=i, column=j, value=val)
                if i == 3:
                    cell.font = Font(bold=True, color="ffffff")
                    cell.fill = PatternFill(start_color="1565c0", end_color="1565c0", fill_type="solid")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    if j == 1:
                        cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center")
        
        ws_stats.column_dimensions['A'].width = 25
        ws_stats.column_dimensions['B'].width = 15
        
        # ===================== FEUILLE 3 : DÉTAIL DES SCORES =====================
        ws_scores = wb.create_sheet("Détail des scores")
        
        # En-têtes détaillés
        score_headers = ["Rang", "N° Dossier", "Nom", "Prénom", "Score", "Décision", "Statut"]
        score_col_widths = [5, 12, 18, 18, 8, 22, 14]
        
        for col, header in enumerate(score_headers, 1):
            cell = ws_scores.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, color="ffffff")
            cell.fill = PatternFill(start_color="0d47a1", end_color="0d47a1", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws_scores.column_dimensions[get_column_letter(col)].width = score_col_widths[col-1]
        
        # Données des scores
        row_idx = 2
        for rank, c in enumerate(sorted_candidats, 1):
            score = float(c.get('score', 0))
            ws_scores.cell(row=row_idx, column=1, value=rank).alignment = number_alignment
            ws_scores.cell(row=row_idx, column=2, value=c.get('numero_dossier', '')).alignment = center_alignment
            ws_scores.cell(row=row_idx, column=3, value=c.get('nom', '')).alignment = cell_alignment
            ws_scores.cell(row=row_idx, column=4, value=c.get('prenom', '')).alignment = cell_alignment
            ws_scores.cell(row=row_idx, column=5, value=score).alignment = number_alignment
            
            # Couleur du score
            if score >= 11:
                color_fill = PatternFill(start_color="c8e6c9", end_color="c8e6c9", fill_type="solid")
            elif score >= 7:
                color_fill = PatternFill(start_color="fff9c4", end_color="fff9c4", fill_type="solid")
            else:
                color_fill = PatternFill(start_color="ffcdd2", end_color="ffcdd2", fill_type="solid")
            
            for col in range(1, 8):
                ws_scores.cell(row=row_idx, column=col).fill = color_fill
            
            ws_scores.cell(row=row_idx, column=6, value=c.get('decision', '')).alignment = cell_alignment
            ws_scores.cell(row=row_idx, column=7, value=c.get('statut', 'en_attente').upper()).alignment = center_alignment
            row_idx += 1
        
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf
    except Exception as e:
        logger.error(f"Erreur generation Excel: {e}")
        import traceback
        traceback.print_exc()
        return None

def generate_pdf_report_enhanced(candidats, poste_filter=""):
    if not REPORTLAB_AVAILABLE:
        return None
    try:
        sorted_candidats = sort_candidats(candidats)
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=1.0*cm, leftMargin=1.0*cm, topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        
        # Styles modernes
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Title'],
            fontSize=16,
            textColor=colors.HexColor('#0d47a1'),
            alignment=TA_CENTER,
            spaceAfter=4,
            fontName='Helvetica-Bold'
        )
        subtitle_style = ParagraphStyle(
            'CustomSubtitle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#666666'),
            alignment=TA_CENTER,
            spaceAfter=8
        )
        header_style = ParagraphStyle(
            'CustomHeader',
            parent=styles['Normal'],
            fontSize=7,
            textColor=colors.whitesmoke,
            alignment=TA_CENTER,
            backColor=colors.HexColor('#0d47a1'),
            fontName='Helvetica-Bold'
        )
        cell_style = ParagraphStyle(
            'CustomCell',
            parent=styles['Normal'],
            fontSize=6,
            alignment=TA_LEFT,
            leading=8
        )
        cell_center = ParagraphStyle(
            'CustomCellCenter',
            parent=styles['Normal'],
            fontSize=6,
            alignment=TA_CENTER,
            leading=8
        )
        stats_style = ParagraphStyle(
            'StatsStyle',
            parent=styles['Normal'],
            fontSize=9,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#1a3a5c'),
            fontName='Helvetica-Bold'
        )
        
        story = []
        story.append(Paragraph("📊 RAPPORT DE CANDIDATURES", title_style))
        story.append(Paragraph(f"Poste: {poste_filter if poste_filter else 'Tous les postes'} | Genere le {datetime.datetime.now().strftime('%d/%m/%Y a %H:%M')}", subtitle_style))
        story.append(Spacer(1, 0.2*cm))
        
        # Statistiques
        total = len(sorted_candidats)
        retenus = sum(1 for c in sorted_candidats if c.get('statut') == 'retenu')
        entretien = sum(1 for c in sorted_candidats if c.get('statut') == 'entretien')
        rejetes = sum(1 for c in sorted_candidats if c.get('statut') == 'rejete')
        en_attente = sum(1 for c in sorted_candidats if c.get('statut') == 'en_attente')
        
        stats_text = f"<b>Total:</b> {total}   |   <b>Retenus:</b> {retenus}   |   <b>Entretien:</b> {entretien}   |   <b>En attente:</b> {en_attente}   |   <b>Rejetés:</b> {rejetes}"
        story.append(Paragraph(stats_text, stats_style))
        story.append(Spacer(1, 0.3*cm))
        
        if not sorted_candidats:
            story.append(Paragraph("Aucun candidat trouve.", styles['Normal']))
        else:
            # Tableau avec Rang en première colonne
            data = [
                [
                    Paragraph("<b>Rang</b>", header_style),
                    Paragraph("<b>N°</b>", header_style),
                    Paragraph("<b>Nom</b>", header_style),
                    Paragraph("<b>Prenom</b>", header_style),
                    Paragraph("<b>Poste</b>", header_style),
                    Paragraph("<b>Statut</b>", header_style),
                    Paragraph("<b>Score</b>", header_style),
                    Paragraph("<b>Motifs</b>", header_style)
                ]
            ]
            
            rank = 1
            for c in sorted_candidats[:200]:
                statut = c.get('statut', 'en_attente')
                try:
                    score = float(c.get('score', 0))
                except (ValueError, TypeError):
                    score = 0.0
                
                # Récupération des motifs
                analyse_details = c.get('analyse_details_parsed', {})
                flags_elim = analyse_details.get('flags_eliminatoires', [])
                points_vigilance = analyse_details.get('points_vigilance', [])
                points_forts = analyse_details.get('points_forts', [])
                synthese = analyse_details.get('synthese_recruteur', '')
                
                # Construction du motif complet
                motif_parts = []
                if flags_elim:
                    for f in flags_elim[:3]:
                        motif_parts.append(f"🚫 {f}")
                if points_vigilance:
                    for p in points_vigilance[:3]:
                        motif_parts.append(f"⚠️ {p}")
                if points_forts:
                    for p in points_forts[:2]:
                        motif_parts.append(f"✅ {p}")
                if synthese:
                    motif_parts.append(f"📝 {synthese[:100]}")
                
                motif = " | ".join(motif_parts) if motif_parts else "Analyse en cours..."
                
                # Statut avec couleur
                if statut == 'retenu':
                    statut_text = f'<font color="#16a34a"><b>✅ Retenu</b></font>'
                elif statut == 'entretien':
                    statut_text = f'<font color="#d97706"><b>📅 Entretien</b></font>'
                elif statut == 'rejete':
                    statut_text = f'<font color="#dc2626"><b>❌ Rejeté</b></font>'
                else:
                    statut_text = f'<font color="#d97706">⏳ En attente</font>'
                
                # Couleur du score
                if score >= 11:
                    score_color = "#16a34a"
                elif score >= 7:
                    score_color = "#d97706"
                else:
                    score_color = "#dc2626"
                
                data.append([
                    Paragraph(str(rank), cell_center),
                    Paragraph(c.get('numero_dossier', ''), cell_center),
                    Paragraph(c.get('nom', '')[:15], cell_style),
                    Paragraph(c.get('prenom', '')[:15], cell_style),
                    Paragraph(c.get('poste', '')[:20], cell_style),
                    Paragraph(statut_text, cell_center),
                    Paragraph(f'<font color="{score_color}"><b>{score:.1f}</b></font>', cell_center),
                    Paragraph(motif[:120], cell_style)
                ])
                rank += 1
            
            # Largeurs des colonnes ajustées pour tout contenir
            col_widths = [1.2*cm, 1.8*cm, 2.5*cm, 2.5*cm, 4.0*cm, 2.5*cm, 1.5*cm, 5.5*cm]
            table = Table(data, colWidths=col_widths, repeatRows=1)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d47a1')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 7),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('VALIGN', (0, 0), (-1, 0), 'MIDDLE'),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
                ('TOPPADDING', (0, 0), (-1, 0), 4),
                ('GRID', (0, 0), (-1, 0), 0.5, colors.HexColor('#333333')),
                ('FONTSIZE', (0, 1), (-1, -1), 6),
                ('VALIGN', (0, 1), (-1, -1), 'MIDDLE'),
                ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
                ('ALIGN', (0, 1), (0, -1), 'CENTER'),
                ('ALIGN', (1, 1), (1, -1), 'CENTER'),
                ('ALIGN', (5, 1), (6, -1), 'CENTER'),
                ('GRID', (0, 1), (-1, -1), 0.3, colors.HexColor('#CCCCCC')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8F9FA')]),
            ]))
            story.append(table)
            story.append(Spacer(1, 0.3*cm))
            story.append(Paragraph(f"<b>Total:</b> {len(sorted_candidats)} candidat(s)", styles['Normal']))
        
        doc.build(story)
        buf.seek(0)
        return buf
    except Exception as e:
        logger.error(f"Erreur generation PDF: {e}")
        import traceback
        traceback.print_exc()
        return None

def generate_csv_report_enhanced(candidats, poste_filter=""):
    sorted_candidats = sort_candidats(candidats)
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    headers = ["Rang", "N° Dossier", "Nom", "Prénom", "Email", "Téléphone", "Poste", "Statut", "Score", "Décision", "Motifs"]
    writer.writerow(headers)
    rank = 1
    for c in sorted_candidats:
        analyse_details = c.get('analyse_details_parsed', {})
        flags_elim = analyse_details.get('flags_eliminatoires', [])
        points_vigilance = analyse_details.get('points_vigilance', [])
        points_forts = analyse_details.get('points_forts', [])
        synthese = analyse_details.get('synthese_recruteur', '')
        statut = c.get('statut', 'en_attente')
        
        # Construction du motif
        motif_parts = []
        if flags_elim:
            motif_parts.append(f"Eliminatoire: {', '.join(flags_elim[:2])}")
        if points_vigilance:
            motif_parts.append(f"Vigilance: {', '.join(points_vigilance[:2])}")
        if points_forts:
            motif_parts.append(f"Forces: {', '.join(points_forts[:2])}")
        if synthese:
            motif_parts.append(f"Synthèse: {synthese[:100]}")
        
        motif = " | ".join(motif_parts) if motif_parts else "Analyse en cours..."
        
        writer.writerow([
            rank,
            c.get('numero_dossier', ''),
            c.get('nom', ''),
            c.get('prenom', ''),
            c.get('email', ''),
            c.get('telephone', ''),
            c.get('poste', ''),
            statut,
            c.get('score', 0),
            c.get('decision', ''),
            motif
        ])
        rank += 1
    return output.getvalue()

def generate_word_report_enhanced(candidats, poste_filter=""):
    if not DOCX_AVAILABLE:
        return None
    try:
        sorted_candidats = sort_candidats(candidats)
        from docx import Document as DocxDocument
        from docx.shared import Inches, Pt, Cm, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        doc = DocxDocument()
        
        def set_cell_border(cell, **kwargs):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for edge in ['top', 'left', 'bottom', 'right']:
                tag = f'w:{edge}'
                if edge in kwargs:
                    border = OxmlElement(tag)
                    border.set(qn('w:val'), 'single')
                    border.set(qn('w:sz'), '4')
                    border.set(qn('w:space'), '0')
                    border.set(qn('w:color'), '1565C0')
                    tcPr.append(border)
        
        # Titre
        title = doc.add_heading(f"📊 RAPPORT DE CANDIDATURES", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Poste: {poste_filter if poste_filter else 'Tous les postes'} | Genere le {datetime.datetime.now().strftime('%d/%m/%Y a %H:%M')}")
        doc.add_paragraph()
        
        # Statistiques
        total = len(sorted_candidats)
        retenus = sum(1 for c in sorted_candidats if c.get('statut') == 'retenu')
        entretien = sum(1 for c in sorted_candidats if c.get('statut') == 'entretien')
        rejetes = sum(1 for c in sorted_candidats if c.get('statut') == 'rejete')
        
        p = doc.add_paragraph()
        p.add_run(f"Total: {total} | ").bold = True
        p.add_run(f"Retenus: {retenus} | ")
        p.add_run(f"Entretien: {entretien} | ")
        p.add_run(f"Rejetés: {rejetes}")
        doc.add_paragraph()
        
        if not sorted_candidats:
            doc.add_paragraph("Aucun candidat trouve.")
        else:
            doc.add_heading("Liste des candidats", level=1)
            # Tableau avec Rang en première colonne
            table = doc.add_table(rows=1, cols=9)
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            header_cells = table.rows[0].cells
            headers = ["Rang", "N° Dossier", "Nom", "Prénom", "Poste", "Statut", "Score", "Décision", "Motifs"]
            for i, header in enumerate(headers):
                header_cells[i].text = header
                header_cells[i].paragraphs[0].runs[0].font.bold = True
                header_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
                header_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'solid')
                shading.set(qn('w:color'), 'auto')
                shading.set(qn('w:fill'), '0d47a1')
                header_cells[i]._tc.get_or_add_tcPr().append(shading)
            
            rank = 1
            for c in sorted_candidats[:200]:
                row_cells = table.add_row().cells
                analyse_details = c.get('analyse_details_parsed', {})
                flags_elim = analyse_details.get('flags_eliminatoires', [])
                points_vigilance = analyse_details.get('points_vigilance', [])
                points_forts = analyse_details.get('points_forts', [])
                synthese = analyse_details.get('synthese_recruteur', '')
                statut = c.get('statut', 'en_attente')
                
                # Construction du motif
                motif_parts = []
                if flags_elim:
                    motif_parts.append(f"Eliminatoire: {', '.join(flags_elim[:2])}")
                if points_vigilance:
                    motif_parts.append(f"Vigilance: {', '.join(points_vigilance[:2])}")
                if points_forts:
                    motif_parts.append(f"Forces: {', '.join(points_forts[:2])}")
                if synthese:
                    motif_parts.append(f"Synthèse: {synthese[:100]}")
                
                motif = " | ".join(motif_parts) if motif_parts else "Analyse en cours..."
                
                row_cells[0].text = str(rank)
                row_cells[1].text = c.get('numero_dossier', '')
                row_cells[2].text = c.get('nom', '')
                row_cells[3].text = c.get('prenom', '')
                row_cells[4].text = c.get('poste', '')[:35]
                row_cells[5].text = statut.upper()
                row_cells[6].text = str(c.get('score', 0))
                row_cells[7].text = c.get('decision', '')[:30]
                row_cells[8].text = motif[:70]
                
                # Alignements
                for idx in [0, 1, 5, 6]:
                    row_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for idx in [2, 3, 4, 7, 8]:
                    row_cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                # Couleur de fond selon statut
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'solid')
                shading.set(qn('w:color'), 'auto')
                if statut == 'retenu':
                    shading.set(qn('w:fill'), 'E8F5E9')
                elif statut == 'entretien':
                    shading.set(qn('w:fill'), 'FFF3E0')
                elif statut == 'rejete':
                    shading.set(qn('w:fill'), 'FFEBEE')
                else:
                    shading.set(qn('w:fill'), 'F5F5F5')
                for cell in row_cells:
                    cell._tc.get_or_add_tcPr().append(shading)
                
                rank += 1
            
            for i, width in enumerate([1.0, 2.2, 3.2, 3.2, 4.0, 2.5, 1.5, 3.0, 4.5]):
                table.columns[i].width = Cm(width)
            
            doc.add_paragraph()
            doc.add_paragraph(f"Total: {len(sorted_candidats)} candidat(s)")
        
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf
    except Exception as e:
        logger.error(f"Erreur generation Word: {e}")
        return None
